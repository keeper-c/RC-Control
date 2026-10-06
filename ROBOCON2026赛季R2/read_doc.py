from docx import Document
import json

doc = Document(r"C:\Users\zzx20\Desktop\RC-R2_机器人控制代码架构文档_V3.3.docx")

print(f"总段落数: {len(doc.paragraphs)}")
print(f"总表格数: {len(doc.tables)}")
print()

# Print all paragraphs with their style
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"P{i} [{p.style.name}] {text[:120]}")
    else:
        print(f"P{i} [{p.style.name}] (空行)")

print("\n=== 表格 ===")
for ti, table in enumerate(doc.tables):
    print(f"\nTable {ti}: {len(table.rows)}行 x {len(table.columns)}列")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip()[:50] for cell in row.cells]
        print(f"  行{ri}: {cells}")
